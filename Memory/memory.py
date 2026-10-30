"""Search and maintain the local Excel history without importing browser workflows."""

from __future__ import annotations

import argparse
import ast
import hashlib
import json
import os
from pathlib import Path
import re
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from Backend.config import lebanon_now

DEFAULT_BOOK = ROOT / "Memory" / "airtable_memory.xlsx"
HEADERS = {
    "Workflows": ("ID", "Function", "Category", "Source", "Summary", "References", "Verification", "Evidence IDs", "Notes"),
    "Functions": ("ID", "Function", "Category", "Source", "Line", "Summary", "References", "Used by", "State", "Verification", "Evidence IDs", "Notes", "Code fingerprint", "Context fingerprint"),
    "Patterns": ("ID", "Topic", "Scope", "Observation", "Reuse check", "Related functions", "Evidence IDs", "Status", "Verified at", "Updated", "Notes", "Code snapshot"),
    "Failures & fixes": ("ID", "Topic", "Scope", "Symptom", "Cause", "Fix", "Verification", "Related functions", "Evidence IDs", "Status", "Verified at", "Updated", "Notes", "Code snapshot"),
    "Evidence": ("ID", "Kind", "Source", "Scope", "Observed result", "Observed at", "Limitation", "Related functions"),
}
# Retain the former status so historical rows and user notes remain valid.
STATUSES = ("Code observed", "Recording observed", "Inspection observed", "Conversation reported", "Investigating", "Live verified", "Needs revalidation")
FUNCTIONS = (ast.FunctionDef, ast.AsyncFunctionDef)


# Function IDs survive source moves; new functions keep their actual source IDs.
FUNCTION_MOVES = {
    'Backend/airtable_actions.py': ('backend/airtable_actions.py', frozenset((
        'ActionReport _new_report _register _registry _require_one action_writes_data '
        'failed_action_report list_actions run_action '
    ).split())),
    'Backend/browser.py': ('backend/browser.py', frozenset((
        'AsyncBrowserSession BrowserSession _browser_options _configure_context '
        '_screenshot_path '
    ).split())),
    'Backend/config.py': ('backend/config.py', frozenset((
        'Settings _LebanonFormatter _env_bool _env_path _env_positive_int _load_local_env '
        '_validate_airtable_url artifact_path configure_logging lebanon_now load_apps '
        'resolve_apps_inventory '
    ).split())),
    'Backend/helpers.py': ('backend/helpers.py', frozenset((
        '_add_root_condition _canonical _click _click_option _close_visibility _groups_named '
        '_one _open_activation_qa_details _record_container _run_validated_section '
        '_stable_condition_row _text _validation_only _visible _wait_for_interface_save '
        '_wait_until _workflow_failure '
    ).split())),
    'Backend/old_core_global_changes.py': ('backend/old_core_global_changes.py', frozenset((
        '_apply_change_request_filters _blank_change_request_rows _brief_details_is_open '
        '_change_request_add_button _change_request_condition_exists _change_request_pattern '
        '_click_smart_import_field _close_fields_popover _configure_change_request_condition '
        '_delete_visibility_is_correct _ensure_delete_visibility _ensure_inline_editing '
        '_ensure_smart_import_table _ensure_suggested_metric_hidden _inline_editing_toggle '
        '_open_change_request_filters _open_delete_button_settings _open_delete_visibility '
        '_open_flow_core_brief_details _open_flow_core_planning _open_media_plan_fields '
        '_page_table_is_added _paragraph_texts _planning_list_is_selected '
        '_remove_signature_from_editor _run_validated_section _smart_import_field_is_configured '
        '_smart_import_option_text_matches _smart_import_table _suggested_metric_is_hidden '
        '_suggested_metric_state _unchanged _wait_for_email_editors _wait_for_interface_save '
        '_wait_for_selected_action _wait_for_smart_import_frame change_request_add_filter '
        'hide_delete_button hide_suggested_metric remove_publicis_media_team_signature '
        'smart_import_core '
    ).split())),
    'Backend/recorder.py': ('backend/recorder.py', frozenset((
        'ClickRecorder _path install_recorder wait_for_recording '
    ).split())),
    'Backend/reporting.py': ('backend/reporting.py', frozenset((
        '_failure_information _interface_warning_sheet _message generate_excel_report '
    ).split())),
    'Backend/runner.py': ('backend/runner.py', frozenset((
        'ConcurrentRunner RunSummary '
    ).split())),
    'Backend/workflows/__init__.py': ('functions.py', frozenset((
        'Assign_Change_Activator change_request_add_filter configure_activation_assignment '
        'configure_activation_automation_scripts configure_activation_qa_visibility '
        'hide_suggestion_details push_master_apps '
    ).split())),
    'Backend/workflows/activation_assignment.py': ('backend/helpers.py', frozenset((
        '_add_group_condition_at _apply_button_visibility _button_visibility_state _click_group '
        '_click_sidebar_heading _condition_labels _configure_condition _configure_interface '
        '_ensure_change_button _ensure_plain_checklist _ensure_qa_checklist _filter_labels '
        '_flat_state _has_condition _nested_group_labels _numbered_condition_row '
        '_open_activation_details _open_visibility _plain_checklist_complete '
        '_plain_checklist_state _plain_conditions_complete _qa_checklist_state '
        '_select_change_activator _select_checklist_group _set_condition_connector '
        '_validate_button_visibility _validate_change_button _validate_plain_checklist '
        '_validate_qa_checklist build_activation_assignment_workflow '
    ).split())),
    'Backend/workflows/activation_qa_visibility.py': ('backend/helpers.py', frozenset((
        '_activation_qa_condition_is_correct _apply_activation_qa_visibility '
        '_open_activation_qa_visibility _resolve_activation_qa_cell '
        '_validate_activation_qa_visibility build_activation_qa_visibility_workflow '
    ).split())),
    'Backend/workflows/activation_scripts.py': ('backend/helpers.py', frozenset((
        '_add_assignment_script _add_reset_script _apply_assignment_automation '
        '_assignment_outer_group _automation_enabled _configure_assignment_automation '
        '_configure_reset_automation _configure_script_variable _delete_assignment_update '
        '_enable_reset_automation _ensure_script_action _finish_script_editor _open_automation '
        '_open_automations _open_existing_script _run_script_summaries _script_dialog '
        '_script_input_row _script_variable_is_correct _validate_assignment_automation '
        '_wait_for_automation_save _workflow_action _write_script '
        'build_activation_automation_scripts_workflow '
    ).split())),
    'Backend/workflows/assign_change_activator.py': ('backend/helpers.py', frozenset((
        '_apply_assign_change_visibility _assign_change_visibility_is_correct '
        '_open_assign_change_visibility _open_change_request_details '
        '_wait_for_assign_change_save build_assign_change_activator_workflow '
    ).split())),
    'Backend/workflows/push_master_apps.py': ('backend/helpers.py', frozenset((
        '_click_master_select_all _master_click _master_installation_rows _master_installations '
        '_master_ready _master_table_state _master_wait _open_master_publish_section '
        '_publish_master_interfaces _publish_master_library _push_master_library '
        'build_push_master_apps_workflow '
    ).split())),
    'Backend/workflows/suggestion_details.py': ('backend/helpers.py', frozenset((
        '_ensure_suggestion_details_hidden _open_flow_core_brief_details '
        '_open_suggestion_details_fields _suggestion_details_is_hidden '
        '_suggestion_details_state _visualization_selection_control '
        '_wait_for_brief_detail_sections build_hide_suggestion_details_workflow '
    ).split())),
    'Airtable_RPA.pyw': ('launcher.pyw', frozenset((
        'ActivityWave GlassPanel Launcher build_arguments load_inventory readable_event '
    ).split())),
    'Backend/cli.py': ('main.py', frozenset((
        '_open_and_inspect _parser _positive_int _run _select_apps main '
    ).split())),
}


# Precise locations preserve IDs when functions from several sources share a file.
FUNCTION_RELOCATIONS = {
    ("Backend/cli.py", "_record_workflow"): "main.py::_open_and_inspect",
    ("Backend/helpers.py", "ActionReport.ok"): "backend/airtable_actions.py::ActionReport.ok",
    ("Backend/helpers.py", "ActionReport.has_warnings"): "backend/airtable_actions.py::ActionReport.has_warnings",
    ("Backend/helpers.py", "ActionReport.status"): "backend/airtable_actions.py::ActionReport.status",
    ("Backend/helpers.py", "ActionReport.finish"): "backend/airtable_actions.py::ActionReport.finish",
    ("Backend/helpers.py", "ActionReport.to_dict"): "backend/airtable_actions.py::ActionReport.to_dict",
    ("Backend/helpers.py", "_new_report"): "backend/airtable_actions.py::_new_report",
    ("Backend/helpers.py", "failed_action_report"): "backend/airtable_actions.py::failed_action_report",
    ("Backend/helpers.py", "_require_one"): "backend/airtable_actions.py::_require_one",
    ("Backend/workflows/smart_import_core.py", "_wait_for_smart_import_frame"): "backend/old_core_global_changes.py::_wait_for_smart_import_frame",
    ("Backend/workflows/smart_import_core.py", "_smart_import_option_text_matches"): "backend/old_core_global_changes.py::_smart_import_option_text_matches",
    ("Backend/workflows/smart_import_core.py", "_smart_import_field_is_configured"): "backend/old_core_global_changes.py::_smart_import_field_is_configured",
    ("Backend/workflows/smart_import_core.py", "_click_smart_import_field"): "backend/old_core_global_changes.py::_click_smart_import_field",
    ("Backend/workflows/smart_import_core.py", "_page_table_is_added"): "backend/old_core_global_changes.py::_page_table_is_added",
    ("Backend/helpers.py", "_wait_for_core_interface_save"): "backend/old_core_global_changes.py::_wait_for_interface_save",
    ("Backend/workflows/smart_import_core.py", "_smart_import_table"): "backend/old_core_global_changes.py::_smart_import_table",
    ("Backend/workflows/smart_import_core.py", "_ensure_smart_import_table"): "backend/old_core_global_changes.py::_ensure_smart_import_table",
    ("Backend/workflows/smart_import_core.py", "_inline_editing_toggle"): "backend/old_core_global_changes.py::_inline_editing_toggle",
    ("Backend/workflows/smart_import_core.py", "_ensure_inline_editing"): "backend/old_core_global_changes.py::_ensure_inline_editing",
    ("Backend/helpers.py", "_run_core_validated_section"): "backend/old_core_global_changes.py::_run_validated_section",
    ("Backend/workflows/remove_publicis_media_team_signature.py", "_wait_for_email_editors"): "backend/old_core_global_changes.py::_wait_for_email_editors",
    ("Backend/workflows/remove_publicis_media_team_signature.py", "_wait_for_selected_action"): "backend/old_core_global_changes.py::_wait_for_selected_action",
    ("Backend/workflows/remove_publicis_media_team_signature.py", "_paragraph_texts"): "backend/old_core_global_changes.py::_paragraph_texts",
    ("Backend/workflows/remove_publicis_media_team_signature.py", "_remove_signature_from_editor"): "backend/old_core_global_changes.py::_remove_signature_from_editor",
    ("Backend/workflows/smart_import_core.py", "smart_import_core"): "backend/old_core_global_changes.py::smart_import_core",
    ("Backend/workflows/remove_publicis_media_team_signature.py", "remove_publicis_media_team_signature"): "backend/old_core_global_changes.py::remove_publicis_media_team_signature",
    ("Backend/workflows/hide_suggested_metric.py", "_open_flow_core_planning"): "backend/old_core_global_changes.py::_open_flow_core_planning",
    ("Backend/workflows/hide_suggested_metric.py", "_planning_list_is_selected"): "backend/old_core_global_changes.py::_planning_list_is_selected",
    ("Backend/workflows/hide_suggested_metric.py", "_open_flow_core_brief_details"): "backend/old_core_global_changes.py::_open_flow_core_brief_details",
    ("Backend/workflows/hide_suggested_metric.py", "_brief_details_is_open"): "backend/old_core_global_changes.py::_brief_details_is_open",
    ("Backend/workflows/hide_suggested_metric.py", "_close_fields_popover"): "backend/old_core_global_changes.py::_close_fields_popover",
    ("Backend/workflows/hide_suggested_metric.py", "_open_media_plan_fields"): "backend/old_core_global_changes.py::_open_media_plan_fields",
    ("Backend/workflows/hide_suggested_metric.py", "_suggested_metric_state"): "backend/old_core_global_changes.py::_suggested_metric_state",
    ("Backend/workflows/hide_suggested_metric.py", "_ensure_suggested_metric_hidden"): "backend/old_core_global_changes.py::_ensure_suggested_metric_hidden",
    ("Backend/workflows/hide_suggested_metric.py", "_suggested_metric_is_hidden"): "backend/old_core_global_changes.py::_suggested_metric_is_hidden",
    ("Backend/workflows/hide_suggested_metric.py", "hide_suggested_metric"): "backend/old_core_global_changes.py::hide_suggested_metric",
    ("Backend/workflows/hide_suggested_metric.py", "_unchanged"): "backend/old_core_global_changes.py::_unchanged",
    ("Backend/workflows/hide_delete_button.py", "_open_delete_button_settings"): "backend/old_core_global_changes.py::_open_delete_button_settings",
    ("Backend/workflows/hide_delete_button.py", "_open_delete_visibility"): "backend/old_core_global_changes.py::_open_delete_visibility",
    ("Backend/workflows/hide_delete_button.py", "_delete_visibility_is_correct"): "backend/old_core_global_changes.py::_delete_visibility_is_correct",
    ("Backend/workflows/hide_delete_button.py", "_ensure_delete_visibility"): "backend/old_core_global_changes.py::_ensure_delete_visibility",
    ("Backend/workflows/hide_delete_button.py", "hide_delete_button"): "backend/old_core_global_changes.py::hide_delete_button",
    ("Backend/workflows/change_request_add_filter.py", "_open_change_request_filters"): "backend/old_core_global_changes.py::_open_change_request_filters",
    ("Backend/workflows/change_request_add_filter.py", "_change_request_add_button"): "backend/old_core_global_changes.py::_change_request_add_button",
    ("Backend/workflows/change_request_add_filter.py", "_change_request_pattern"): "backend/old_core_global_changes.py::_change_request_pattern",
    ("Backend/workflows/change_request_add_filter.py", "_change_request_condition_exists"): "backend/old_core_global_changes.py::_change_request_condition_exists",
    ("Backend/workflows/change_request_add_filter.py", "_blank_change_request_rows"): "backend/old_core_global_changes.py::_blank_change_request_rows",
    ("Backend/workflows/change_request_add_filter.py", "_configure_change_request_condition"): "backend/old_core_global_changes.py::_configure_change_request_condition",
    ("Backend/workflows/change_request_add_filter.py", "_apply_change_request_filters"): "backend/old_core_global_changes.py::_apply_change_request_filters",
    ("Backend/workflows/change_request_add_filter.py", "change_request_add_filter"): "functions.py::change_request_add_filter",
    ("Backend/workflows/activation_assignment.py", "configure_activation_assignment"): "functions.py::configure_activation_assignment",
    ("Backend/workflows/activation_qa_visibility.py", "configure_activation_qa_visibility"): "functions.py::configure_activation_qa_visibility",
    ("Backend/helpers.py", "_open_automation"): "backend/helpers.py::_open_automation",
    ("Backend/helpers.py", "_workflow_action"): "backend/helpers.py::_workflow_action",
    ("Backend/helpers.py", "_run_script_summaries"): "backend/helpers.py::_run_script_summaries",
    ("Backend/helpers.py", "_assignment_outer_group"): "backend/helpers.py::_assignment_outer_group",
    ("Backend/helpers.py", "_delete_assignment_update"): "backend/helpers.py::_delete_assignment_update",
    ("Backend/helpers.py", "_delete_assignment_update.deleted"): "backend/helpers.py::_delete_assignment_update.deleted",
    ("Backend/helpers.py", "_add_assignment_script"): "backend/helpers.py::_add_assignment_script",
    ("Backend/helpers.py", "_add_reset_script"): "backend/helpers.py::_add_reset_script",
    ("Backend/helpers.py", "_script_dialog"): "backend/helpers.py::_script_dialog",
    ("Backend/helpers.py", "_open_existing_script"): "backend/helpers.py::_open_existing_script",
    ("Backend/helpers.py", "_write_script"): "backend/helpers.py::_write_script",
    ("Backend/helpers.py", "_script_input_row"): "backend/helpers.py::_script_input_row",
    ("Backend/helpers.py", "_script_variable_is_correct"): "backend/helpers.py::_script_variable_is_correct",
    ("Backend/helpers.py", "_configure_script_variable"): "backend/helpers.py::_configure_script_variable",
    ("Backend/helpers.py", "_finish_script_editor"): "backend/helpers.py::_finish_script_editor",
    ("Backend/helpers.py", "_finish_script_editor.closed"): "backend/helpers.py::_finish_script_editor.closed",
    ("Backend/helpers.py", "_ensure_script_action"): "backend/helpers.py::_ensure_script_action",
    ("Backend/helpers.py", "_wait_for_automation_save"): "backend/helpers.py::_wait_for_automation_save",
    ("Backend/helpers.py", "_wait_for_automation_save.saved"): "backend/helpers.py::_wait_for_automation_save.saved",
    ("Backend/helpers.py", "_apply_assignment_automation"): "backend/helpers.py::_apply_assignment_automation",
    ("Backend/helpers.py", "_apply_assignment_automation.add_script"): "backend/helpers.py::_apply_assignment_automation.add_script",
    ("Backend/helpers.py", "_validate_assignment_automation"): "backend/helpers.py::_validate_assignment_automation",
    ("Backend/helpers.py", "_configure_assignment_automation"): "backend/helpers.py::_configure_assignment_automation",
    ("Backend/helpers.py", "_configure_assignment_automation.apply"): "backend/helpers.py::_configure_assignment_automation.apply",
    ("Backend/helpers.py", "_configure_assignment_automation.validate"): "backend/helpers.py::_configure_assignment_automation.validate",
    ("Backend/helpers.py", "_automation_enabled"): "backend/helpers.py::_automation_enabled",
    ("Backend/helpers.py", "_enable_reset_automation"): "backend/helpers.py::_enable_reset_automation",
    ("Backend/helpers.py", "_enable_reset_automation.enabled"): "backend/helpers.py::_enable_reset_automation.enabled",
    ("Backend/helpers.py", "_configure_reset_automation"): "backend/helpers.py::_configure_reset_automation",
    ("Backend/helpers.py", "_configure_reset_automation.add_script"): "backend/helpers.py::_configure_reset_automation.add_script",
    ("Backend/helpers.py", "_configure_reset_automation.apply"): "backend/helpers.py::_configure_reset_automation.apply",
    ("Backend/helpers.py", "_configure_reset_automation.validate"): "backend/helpers.py::_configure_reset_automation.validate",
    ("Backend/workflows/activation_scripts.py", "configure_activation_automation_scripts"): "functions.py::configure_activation_automation_scripts",
    ("Backend/workflows/adsets_metrics_sync.py", "unsync_adsets_metrics"): "Backend/workflows/__init__.py::unsync_adsets_metrics",
    ("Backend/workflows/assign_change_activator.py", "Assign_Change_Activator"): "functions.py::Assign_Change_Activator",
    ("Backend/workflows/push_master_apps.py", "push_master_apps"): "functions.py::push_master_apps",
    ("Backend/workflows/suggestion_details.py", "hide_suggestion_details"): "functions.py::hide_suggestion_details",
    ("Backend/helpers.py", "list_actions"): "backend/airtable_actions.py::list_actions",
    ("Backend/helpers.py", "action_added_on"): "Backend/airtable_actions.py::action_added_on",
    ("Backend/helpers.py", "run_action"): "backend/airtable_actions.py::run_action",
    ("Backend/helpers.py", "action_writes_data"): "backend/airtable_actions.py::action_writes_data",
}
FUNCTION_MERGES = {
    "backend/old_core_global_changes.py::change_request_add_filter": "functions.py::change_request_add_filter",
    "backend/helpers.py::build_activation_assignment_workflow": "functions.py::configure_activation_assignment",
    "backend/helpers.py::build_activation_qa_visibility_workflow": "functions.py::configure_activation_qa_visibility",
    "backend/helpers.py::build_activation_automation_scripts_workflow": "functions.py::configure_activation_automation_scripts",
    "Backend/workflows/adsets_metrics_sync.py::build_unsync_adsets_metrics_workflow": "Backend/workflows/__init__.py::unsync_adsets_metrics",
    "backend/helpers.py::build_assign_change_activator_workflow": "functions.py::Assign_Change_Activator",
    "backend/helpers.py::build_push_master_apps_workflow": "functions.py::push_master_apps",
    "backend/helpers.py::build_hide_suggestion_details_workflow": "functions.py::hide_suggestion_details",
}


# 1. Build a static catalogue. References are candidates, not an execution trace.
def catalogue(root: Path = ROOT) -> dict[str, dict]:
    paths = sorted([*root.glob("*.py"), *root.glob("*.pyw"), *(root / "Backend").rglob("*.py")])
    modules = {}
    for path in paths:
        parts = list(path.relative_to(root).with_suffix("").parts)
        if parts[-1] == "__init__":
            parts.pop()
        modules[".".join(parts)] = path
    trees, texts, imports, entries = {}, {}, {}, {}
    workflow_targets = set()
    helpers_path = root / "Backend" / "helpers.py"
    if helpers_path.exists():
        for node in ast.parse(helpers_path.read_text(encoding="utf-8-sig")).body:
            if isinstance(node, ast.AnnAssign) and isinstance(node.target, ast.Name) and node.target.id == "_ACTIONS" and isinstance(node.value, ast.Dict):
                for key, value in zip(node.value.keys, node.value.values):
                    if isinstance(key, ast.Constant) and isinstance(value, ast.Call) and value.args and isinstance(value.args[0], ast.Constant):
                        workflow_targets.add((value.args[0].value, key.value))
    for module, path in modules.items():
        source = path.read_text(encoding="utf-8-sig")
        tree = ast.parse(source, filename=str(path))
        trees[module], texts[module] = tree, source
        imports[module] = {}
        for node in tree.body:
            if isinstance(node, ast.ImportFrom):
                imported_module = ("Backend" + node.module[7:]) if node.module and node.module.startswith("backend.") else node.module
                if imported_module in modules:
                    for alias in node.names:
                        imports[module][alias.asname or alias.name] = (imported_module, alias.name)

        def visit(nodes, prefix=""):
            for node in nodes:
                if isinstance(node, (*FUNCTIONS, ast.ClassDef)):
                    name = prefix + node.name
                    if isinstance(node, FUNCTIONS):
                        source_path = path.relative_to(root).as_posix()
                        original_source, moved_names = FUNCTION_MOVES.get(source_path, (source_path, ()))
                        id_source = original_source if name.split(".")[0] in moved_names else source_path
                        key = FUNCTION_RELOCATIONS.get((source_path, name), f"{id_source}::{name}")
                        if key in entries:
                            raise ValueError(f"Duplicate function ID: {key}")
                        category = "Workflow" if (module, name) in workflow_targets else (
                            "Helper" if module == "Backend.helpers" or module.startswith("Backend.workflows.") else "Infrastructure")
                        segment = ast.get_source_segment(source, node) or ""
                        entries[key] = {"ID": key, "Function": name, "Category": category,
                                        "Source": path.relative_to(root).as_posix(), "Line": node.lineno,
                                        "Summary": ast.get_docstring(node) or "", "References": "", "Used by": "",
                                        "State": "Current", "Verification": "No surviving live verification linked",
                                        "Evidence IDs": "", "Notes": "", "Code fingerprint": digest(segment),
                                        "_module": module, "_node": node}
                    visit(node.body, name + ".")
        visit(tree.body)

    lookup = {(e["_module"], e["Function"]): key for key, e in entries.items()}
    for key, entry in entries.items():
        names = set()

        class References(ast.NodeVisitor):
            def visit_Name(self, node):
                if isinstance(node.ctx, ast.Load):
                    names.add(node.id)

            def visit_FunctionDef(self, node):
                if node is entry["_node"]:
                    self.generic_visit(node)

            visit_AsyncFunctionDef = visit_FunctionDef

        References().visit(entry["_node"])
        found = set()
        for name in names:
            parts = entry["Function"].split(".")[:-1]
            options = [(entry["_module"], ".".join(parts[:i] + [name])) for i in range(len(parts), -1, -1)]
            options.append(imports[entry["_module"]].get(name, ("", "")))
            target = next((lookup[option] for option in options if option in lookup), None)
            if target and target != key:
                found.add(target)
        entry["References"] = "\n".join(sorted(found))
        for target in found:
            entries[target]["Used by"] += key + "\n"

    # Conservative fingerprints include local imported modules and constants.
    def context(module):
        pending, seen = [module], set()
        while pending:
            item = pending.pop()
            if item in seen:
                continue
            seen.add(item)
            for node in trees[item].body:
                if isinstance(node, ast.ImportFrom):
                    imported_module = ("Backend" + node.module[7:]) if node.module and node.module.startswith("backend.") else node.module
                    if imported_module in modules:
                        pending.append(imported_module)
        return digest("\n".join(item + ":" + digest(texts[item]) for item in sorted(seen)))

    contexts = {module: context(module) for module in modules}
    for entry in entries.values():
        entry["Context fingerprint"] = contexts[entry.pop("_module")]
        entry["Used by"] = entry["Used by"].strip()
        entry.pop("_node")
    return entries


def digest(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


# 2. Excel is the only durable memory store; searches do not create another database.
def read_records(path: Path) -> dict[str, list[dict]]:
    if not path.exists():
        raise ValueError(f"Memory workbook not found: {path}. Run refresh first.")
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        result = {}
        for name, headers in HEADERS.items():
            if name not in book.sheetnames:
                raise ValueError(f"Missing memory sheet: {name}")
            sheet = book[name]
            actual = tuple(cell.value for cell in sheet[4])
            if actual != headers:
                raise ValueError(f"Do not rename or reorder columns in {name}; row 4 must match the memory schema.")
            rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=5, values_only=True) if row[0]]
            if len({row["ID"] for row in rows}) != len(rows):
                raise ValueError(f"Duplicate IDs in {name}; resolve before updating.")
            result[name] = [{k: "" if v is None else v for k, v in row.items()} for row in rows]
        return result
    finally:
        book.close()


def split_ids(value) -> list[str]:
    return [line.strip() for line in str(value or "").splitlines() if line.strip()]


def snapshot(row, entries) -> str:
    keys = split_ids(row.get("Related functions"))
    return json.dumps({key: entries[key]["Context fingerprint"] if key in entries else "MISSING"
                       for key in sorted(keys)}, sort_keys=True)


def freshness(row, entries) -> str:
    saved = row.get("Code snapshot")
    if not saved:
        return "Untracked code snapshot"
    return "Code changed; needs revalidation" if saved != snapshot(row, entries) else "Code unchanged; check current DOM"


def refresh(records, entries):
    previous = {row["ID"]: row for row in records["Functions"]}
    combined = dict(entries)
    for key, row in entries.items():
        old = previous.get(key, {})
        for field in ("Verification", "Evidence IDs", "Notes"):
            row[field] = old.get(field, row[field])
        if old and old.get("Context fingerprint") != row["Context fingerprint"]:
            row["State"] = "Code changed; review linked verification"
        elif old.get("State", "").startswith("Code changed"):
            row["State"] = old["State"]
    for key in previous.keys() - entries.keys():
        target = entries.get(FUNCTION_MERGES.get(key))
        combined[key] = {**previous[key], "State": "Removed from current source; historical entry"}
        if target:
            combined[key].update(Source=target["Source"], Line=target["Line"],
                                 Category="Historical function", State=f"Merged into {target['ID']}; historical entry")
    records["Functions"] = sorted(combined.values(), key=lambda row: (row["Category"], row["ID"]))
    records["Workflows"] = [{field: row.get(field, "") for field in HEADERS["Workflows"]}
                            for row in records["Functions"] if row["ID"] in entries and row["Category"] == "Workflow"]
    for sheet in ("Patterns", "Failures & fixes"):
        for row in records[sheet]:
            if row.get("Code snapshot") and freshness(row, entries).startswith("Code changed"):
                row["Status"] = "Needs revalidation"
    return records


def save_records(path: Path, records, expected_hash=None):
    book = load_workbook(path) if path.exists() else Workbook()
    for sheet in list(book):
        if sheet.title in {"Start here", *HEADERS} or (sheet.title == "Sheet" and sheet.max_row == 1 and sheet["A1"].value is None):
            book.remove(sheet)
    guide = book.create_sheet("Start here", 0)
    guide.append(["AIRTABLE PROJECT MEMORY", "How to read and maintain this workbook"])
    guide.append(["Updated (Lebanon)", lebanon_now().isoformat(timespec="seconds")])
    for name, explanation in (
        ("Workflows", "Start here: all workflows in the shared catalogue, with linked verification."),
        ("Functions", "Generated function index, references and callers. Edit only Verification, Evidence IDs and Notes; use record for these too."),
        ("Patterns", "Scoped reusable observations. Code observed is not a successful live trial."),
        ("Failures & fixes", "Symptoms, causes, corrections, evidence and verification. Conversation-reported history remains distinct from live proof."),
        ("Evidence", "Pointers and concise findings only; original recordings, logs, reports and profiles remain local."),
        ("Commands", 'python -B Memory/memory.py search "saving before reload"\npython -B Memory/memory.py show <ID>\npython -B Memory/memory.py refresh\npython -B Memory/memory.py record --file <record.json>'),
        ("Verified status", "Live verified requires a specific observed result, scope, verification date and referenced Live trial evidence. A successful skip does not verify a write path."),
        ("Staleness", "Search checks code against stored fingerprints even without refresh. Changes trigger review. Unchanged code does not establish an unchanged Airtable DOM."),
        ("Reference limits", "Python AST references are static candidates, including callbacks. Dynamic dispatch, object method calls and runtime imports may not be resolved."),
        ("Stable IDs", "Existing function IDs survive the repository reorganization. An ID can name an old file; use Source and Line for its current location. New functions use their actual source paths."),
        ("Manual review", "Use filters and wrapped cells. Keep IDs and column names stable. Close Excel before a tool update. Tool saves rebuild managed-sheet styling; extra user sheets and allowed notes survive."),
        ("Update safety", "Atomic replacement protects the existing workbook if saving fails. Other sheets and manual notes are preserved. No automatic history deletion."),
        ("History coverage", "Current source and surviving local artifacts are catalogued. Deleted trials cannot be recreated. Initial records identify their source and limitations."),
        ("Authority", "Memory is a research aid. New user instructions, current recording and verified live state take precedence; memory does not authorize writes, retries or targets."),
    ):
        guide.append([name, explanation])
    for row in guide:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        guide.row_dimensions[row[0].row].height = 48 if row[0].row > 2 else 30
    guide.column_dimensions["A"].width, guide.column_dimensions["B"].width = 28, 115
    guide.freeze_panes = "B3"
    for cell in guide[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True, size=14)
    for row in guide.iter_rows(min_row=3):
        row[0].font = Font(color="17365D", bold=True)
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="EEF5FB" if cell.row % 2 else "FFFFFF")

    for number, (name, headers) in enumerate(HEADERS.items(), 1):
        sheet = book.create_sheet(name)
        sheet.append([name.upper()])
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        sheet.append([f"{len(records[name])} records | Filter row 4 | Read scope and evidence before reuse"])
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        sheet.append([])
        sheet.append(headers)
        for record in records[name]:
            if any(isinstance(value, str) and len(value) > 32767 for value in record.values()):
                raise ValueError(f"Excel text limit exceeded in {name} record {record.get('ID')}; shorten the summary or link its evidence.")
            sheet.append([record.get(field, "") for field in headers])
        if records[name]:
            table = Table(displayName=f"MemoryTable{number}", ref=f"A4:{sheet.cell(sheet.max_row, len(headers)).coordinate}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
        for column, field in enumerate(headers, 1):
            letter = sheet.cell(4, column).column_letter
            sheet.column_dimensions[letter].width = (52 if field in {"Summary", "Observation", "Symptom", "Cause", "Fix", "Reuse check", "Verification", "Observed result", "Notes", "Limitation"}
                                                       else 12 if field == "Line" else 32 if field in {"Topic", "Scope", "Status", "State"} else 45)
            if field in {"Code fingerprint", "Context fingerprint", "Code snapshot"}:
                sheet.column_dimensions[letter].hidden = True
            if field == "Status":
                validation = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"')
                sheet.add_data_validation(validation)
                validation.add(f"{letter}5:{letter}1048576")
        sheet.freeze_panes = "C5"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.print_title_rows = "1:4"
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                # Preserve literal observations; never interpret imported text as formulas.
                if isinstance(cell.value, str):
                    if len(cell.value) > 32767:
                        raise ValueError(f"Excel cell too long: {name}!{cell.coordinate}")
                    cell.data_type = "s"
                if cell.row in {1, 4}:
                    cell.fill = PatternFill("solid", fgColor="17365D")
                    cell.font = Font(color="FFFFFF", bold=True, size=14 if cell.row == 1 else 10)
                elif cell.row >= 5:
                    field = headers[cell.column - 1]
                    if field == "Source" and cell.value:
                        source = str(cell.value).split("#", 1)[0]
                        if (ROOT / source).is_file():
                            cell.hyperlink = os.path.relpath(ROOT / source, path.parent).replace("\\", "/")
                            cell.font = Font(color="0563C1", underline="single")
                    if field in {"Status", "State"}:
                        color = "E2F0D9" if cell.value == "Live verified" else "FFF2CC" if "revalid" in str(cell.value).lower() or "changed" in str(cell.value).lower() else "EEF5FB"
                        cell.fill = PatternFill("solid", fgColor=color)
            sheet.row_dimensions[row[0].row].height = 84 if row[0].row >= 5 else 25
        guide.cell(number + 2, 1).hyperlink = f"#'{name}'!A1"
    destinations = {name: {row["ID"]: index for index, row in enumerate(records[name], 5)} for name in HEADERS}
    for name, headers in HEADERS.items():
        for row in book[name].iter_rows(min_row=5):
            for field, target in (("Evidence IDs", "Evidence"), ("Related functions", "Functions"),
                                  ("References", "Functions"), ("Used by", "Functions")):
                if field not in headers:
                    continue
                cell = row[headers.index(field)]
                linked = split_ids(cell.value)
                if linked and linked[0] in destinations[target]:
                    cell.hyperlink = f"#'{target}'!A{destinations[target][linked[0]]}"
                    cell.font = Font(color="0563C1", underline="single")
    path.parent.mkdir(parents=True, exist_ok=True)
    book.properties.title = "Airtable project knowledge and verified history"
    book.properties.created = lebanon_now().replace(tzinfo=None)
    with tempfile.NamedTemporaryFile(dir=path.parent, prefix=".memory-", suffix=".xlsx", delete=False) as stream:
        temporary = Path(stream.name)
    try:
        book.save(temporary)
        actual_hash = hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else ""
        if expected_hash is not None and actual_hash != expected_hash:
            raise ValueError("Workbook changed during this operation; no changes saved. Retry after reviewing the latest workbook.")
        temporary.replace(path)
    finally:
        book.close()
        temporary.unlink(missing_ok=True)


# 3. Record findings explicitly; never infer a cause or successful fix from logs.
def record_item(records, payload, entries):
    sheet = payload.get("sheet")
    values = payload.get("record", {})
    allowed = {"Patterns", "Failures & fixes", "Evidence", "Functions"}
    if sheet not in allowed or not isinstance(values, dict) or not values.get("ID"):
        raise ValueError("Supply sheet and record with an ID; supported sheets: " + ", ".join(sorted(allowed)))
    if set(values) - set(HEADERS[sheet]):
        raise ValueError("Unknown fields: " + ", ".join(sorted(set(values) - set(HEADERS[sheet]))))
    current = next((row for row in records[sheet] if row["ID"] == values["ID"]), None)
    if sheet == "Functions" and (not current or set(values) - {"ID", "Notes", "Evidence IDs", "Verification"}):
        raise ValueError("Only update Notes, Evidence IDs and Verification on an existing function.")
    merged = {**(current or {}), **values}
    if sheet == "Evidence":
        for field in ("Kind", "Source", "Scope", "Observed result", "Observed at", "Limitation"):
            if not merged.get(field):
                raise ValueError(f"Missing evidence field {field}.")
    if sheet == "Functions" and payload.get("revalidated") and "Verification" in values:
        if merged["ID"] not in entries:
            raise ValueError("Cannot revalidate a function missing from current source.")
        merged["State"] = "Current"
        for field in ("Code fingerprint", "Context fingerprint"):
            merged[field] = entries[merged["ID"]][field]
    if sheet in {"Patterns", "Failures & fixes"}:
        finding_fields = ("Observation", "Reuse check") if sheet == "Patterns" else ("Symptom", "Cause", "Fix", "Verification")
        for field in ("Topic", "Scope", "Related functions", "Evidence IDs", "Status", *finding_fields):
            if not merged.get(field):
                raise ValueError(f"Missing {field}.")
        if merged["Status"] not in STATUSES:
            raise ValueError("Unsupported status.")
        unknown = set(split_ids(merged["Related functions"])) - entries.keys()
        if unknown:
            raise ValueError("Unknown function IDs: " + ", ".join(sorted(unknown)))
        evidence = {row["ID"]: row for row in records["Evidence"]}
        if set(split_ids(merged["Evidence IDs"])) - evidence.keys():
            raise ValueError("Add referenced evidence records first.")
        if merged["Status"] == "Live verified":
            if not merged.get("Verified at") or not any(evidence[key].get("Kind") == "Live trial" for key in split_ids(merged["Evidence IDs"])):
                raise ValueError("Live verified requires Verified at and linked Live trial evidence.")
            proof = merged.get("Verification") if sheet == "Failures & fixes" else merged.get("Reuse check")
            if not proof:
                raise ValueError("Describe the observed verification.")
        # A prose edit cannot silently bless a changed source fingerprint.
        if current and current.get("Code snapshot") != snapshot(current, entries) and not payload.get("revalidated"):
            merged["Status"] = "Needs revalidation"
            merged["Code snapshot"] = current.get("Code snapshot", "")
        else:
            merged["Code snapshot"] = snapshot(merged, entries)
        merged["Updated"] = lebanon_now().isoformat(timespec="seconds")
    if current:
        current.update(merged)
    else:
        records[sheet].append(merged)


def search(records, entries, query: str):
    aliases = {"saving": "save", "saved": "save", "saves": "save", "failures": "failure", "fixes": "fix"}
    ignored = {"a", "an", "the", "to", "of", "in", "on", "and", "or", "for", "with", "before", "after", "when", "is", "was", "be", "it"}
    def tokens(value):
        return {aliases.get(word, word) for word in re.findall(r"[\w]+", str(value).casefold().replace("_", " "))} - ignored
    terms = tokens(query)
    matches = []
    for sheet, rows in records.items():
        if sheet == "Workflows":
            continue
        for row in rows:
            body = " ".join(str(value) for key, value in row.items() if "fingerprint" not in key.lower() and key != "Code snapshot").casefold().replace("_", " ")
            words = tokens(body)
            score = len(terms & words)
            if score:
                title = str(row.get("Topic", row.get("Function", row.get("ID"))))
                score = score * 3 + len(terms & tokens(title)) * 3
                matches.append((score, sheet, row))
    priority = {"Failures & fixes": 0, "Patterns": 1, "Evidence": 2, "Functions": 3}
    return sorted(matches, key=lambda item: (-item[0], priority[item[1]], item[2]["ID"]))


def display(sheet, row, entries):
    result = {"Sheet": sheet, **{key: value for key, value in row.items() if value and key not in {"Code fingerprint", "Context fingerprint", "Code snapshot"}}}
    if sheet in {"Patterns", "Failures & fixes"}:
        result["Current code check"] = freshness(row, entries)
    elif sheet in {"Functions", "Workflows"}:
        live = entries.get(row["ID"])
        result["Current code check"] = ("Source missing" if not live else "Index out of date; run refresh"
                                        if row.get("Context fingerprint", live["Context fingerprint"]) != live["Context fingerprint"] else "Current source indexed; live behavior still requires evidence")
    print(json.dumps(result, ensure_ascii=False, indent=2))


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--book", type=Path, default=DEFAULT_BOOK)
    commands = parser.add_subparsers(dest="command", required=True)
    commands.add_parser("refresh", help="Create/refresh the code catalogue, preserving recorded knowledge")
    find = commands.add_parser("search", help="Search observations, failures and functions")
    find.add_argument("query")
    find.add_argument("--limit", type=int, default=8, help="Maximum displayed matches; use 0 for all")
    show = commands.add_parser("show", help="Read a complete record by its stable ID")
    show.add_argument("id")
    record = commands.add_parser("record", help="Add/update an explicit finding from JSON")
    record.add_argument("--file", type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        entries = catalogue()
        if args.command != "refresh" and not args.book.exists():
            raise ValueError(f"Memory workbook not found: {args.book}. Run refresh first.")
        expected_hash = hashlib.sha256(args.book.read_bytes()).hexdigest() if args.book.exists() else ""
        records = read_records(args.book) if args.book.exists() else {name: [] for name in HEADERS}
        if args.command == "refresh":
            save_records(args.book, refresh(records, entries), expected_hash)
            print(f"Updated {args.book}: {len(records['Workflows'])} workflows; {len(records['Functions'])} functions.")
        elif args.command == "record":
            payload = json.loads(args.file.read_text(encoding="utf-8-sig"))
            for item in payload if isinstance(payload, list) else [payload]:
                record_item(records, item, entries)
            save_records(args.book, refresh(records, entries), expected_hash)
            print(f"Recorded findings in {args.book}.")
        elif args.command == "search":
            matches = search(records, entries, args.query)
            for _, sheet, row in matches[:args.limit] if args.limit > 0 else matches:
                display(sheet, row, entries)
            print(f"{len(matches)} matching records. Use show <ID> for details; --limit 0 displays all matches.")
        else:
            found = [(sheet, row) for sheet, rows in records.items() for row in rows
                     if row["ID"] == args.id and sheet != "Workflows"]
            if not found:
                raise ValueError(f"Record not found: {args.id}")
            for sheet, row in found:
                display(sheet, row, entries)
        return 0
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"Memory update/search failed: {exc}. Close Excel before updates; existing workbook is preserved on save failure.", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
