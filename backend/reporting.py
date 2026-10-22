"""Generate the single user-facing Excel report for an automation run."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from Backend.airtable_actions import Outcome
from Backend.config import artifact_path, lebanon_now
from Backend.runner import AppActionExecution, RunSummary


_NAVY = "17365D"
_LIGHT_BLUE = "EEF5FB"
_GREEN = "E2F0D9"
_YELLOW = "FFF2CC"
_RED = "FCE4D6"
_WHITE = "FFFFFF"
_BORDER = Side(style="thin", color="C9D3DD")


# 1. Readable status and diagnostic text.
def _message(execution: AppActionExecution) -> str:
    return "\n".join(
        f"[{item.outcome.value.upper()}] {item.candidate}: {item.message}"
        for item in sorted(
            execution.report.results,
            key=lambda item: {Outcome.FAILED: 0, Outcome.WARNING: 1}.get(item.outcome, 2),
        )
    )


def _failure_information(execution: AppActionExecution) -> str:
    if execution.report.ok:
        return ""

    information: list[str] = []
    for item in execution.report.results:
        if item.outcome == Outcome.FAILED:
            information.append(f"{item.candidate}: {item.message}")
            if item.details:
                information.append(
                    json.dumps(item.details, ensure_ascii=False, default=str)
                )
    if execution.report.final_url:
        information.append(f"Final URL: {execution.report.final_url}")
    if execution.screenshot_path:
        information.append(f"Screenshot: {execution.screenshot_path}")
    information.append("See terminal output for technical details.")
    return "\n".join(information)


# 2. Explicit non-failing interface warning summary.
def _interface_warning_sheet(workbook: Workbook, summary: RunSummary) -> None:
    rows = []
    for execution in summary.executions:
        for result in execution.report.results:
            if not result.details.get("interface_publication_warning"):
                continue

            def downstream(candidate: str) -> str:
                return next((f"{item.outcome.value.upper()}: {item.message}"
                             for item in execution.report.results if item.candidate == candidate),
                            "Not processed or verified.")

            rows.append((execution.app.name, result.details["interface_error"],
                         downstream("Master app library publication"), downstream("Installed base push"),
                         execution.app.url))
    if not rows:
        return
    sheet = workbook.create_sheet("Interface publish warnings")
    sheet.append(("App continued after interface warning", "Interface publication warning",
                  "Library publication", "Installed base push", "App URL"))
    for row in rows:
        sheet.append(row)
    for cells in sheet:
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=_BORDER)
            if cell.row == 1:
                cell.font = Font(bold=True, color=_WHITE)
                cell.fill = PatternFill("solid", fgColor=_NAVY)
    for column, width in enumerate((36, 68, 68, 68, 55), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


# 3. Formatted workbook and collision-safe output.
def generate_excel_report(summary: RunSummary, output_dir: Path) -> Path:
    """Create exactly one formatted .xlsx workbook for this run."""
    output_dir.mkdir(parents=True, exist_ok=True)
    action_label = "multi-action" if len(summary.actions) > 1 else summary.actions[0]
    safe_action = "".join(
        character if character.isalnum() or character in "-_" else "-"
        for character in action_label
    ).strip("-")
    report_path = artifact_path(output_dir, safe_action, ".xlsx")
    temporary_path = report_path.with_name(f".{report_path.stem}.tmp.xlsx")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Run Report"
    workbook.properties.title = "Airtable Automation Run Report"
    workbook.properties.creator = "Airtable Automation"
    workbook.properties.created = lebanon_now().replace(tzinfo=None)

    worksheet.merge_cells("A1:F1")
    title = worksheet["A1"]
    title.value = "Airtable Automation Run Report"
    title.font = Font(size=16, bold=True, color=_WHITE)
    title.fill = PatternFill("solid", fgColor=_NAVY)
    title.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 28

    worksheet["A2"] = "Started (Lebanon time)"
    worksheet["B2"] = summary.started_at
    worksheet["D2"] = "Finished (Lebanon time)"
    worksheet["E2"] = summary.finished_at
    worksheet["A3"] = "Actions"
    worksheet["B3"] = ", ".join(summary.actions)
    worksheet["D3"] = "Concurrency"
    worksheet["E3"] = summary.concurrency
    worksheet["A4"] = "Overall run result"
    worksheet["B4"] = summary.status
    worksheet["B4"].font = Font(bold=True)
    worksheet["B4"].fill = PatternFill(
        "solid", fgColor=(_YELLOW if summary.has_warnings else _GREEN) if summary.ok else _RED
    )
    for cell_ref in ("A2", "D2", "A3", "D3", "A4"):
        worksheet[cell_ref].font = Font(bold=True, color=_NAVY)

    header_row = 5
    headers = (
        "App",
        "Action",
        "Result",
        "Message",
        "Duration",
        "Failure Information",
    )
    for column, heading in enumerate(headers, start=1):
        cell = worksheet.cell(row=header_row, column=column, value=heading)
        cell.font = Font(bold=True, color=_WHITE)
        cell.fill = PatternFill("solid", fgColor=_NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=_BORDER)

    status_fills = {
        "Success": PatternFill("solid", fgColor=_GREEN),
        "Success with warnings": PatternFill("solid", fgColor=_YELLOW),
        "Skipped": PatternFill("solid", fgColor=_YELLOW),
        "Failed": PatternFill("solid", fgColor=_RED),
    }
    for row_number, execution in enumerate(summary.executions, start=header_row + 1):
        result = execution.report.status
        values = (
            execution.app.name,
            execution.action,
            result,
            _message(execution),
            execution.duration_seconds,
            _failure_information(execution),
        )
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=_BORDER)
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=_LIGHT_BLUE)
        app_cell = worksheet.cell(row=row_number, column=1)
        app_cell.hyperlink = execution.app.url
        app_cell.style = "Hyperlink"
        result_cell = worksheet.cell(row=row_number, column=3)
        result_cell.fill = status_fills[result]
        result_cell.font = Font(bold=True)
        result_cell.alignment = Alignment(horizontal="center", vertical="top")
        worksheet.cell(row=row_number, column=5).number_format = '0.00 "s"'

    widths = (32, 40, 24, 68, 14, 68)
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width
    worksheet.freeze_panes = "A6"
    if summary.executions:
        worksheet.auto_filter.ref = f"A{header_row}:F{header_row + len(summary.executions)}"
    worksheet.sheet_view.showGridLines = False
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.print_title_rows = f"1:{header_row}"

    _interface_warning_sheet(workbook, summary)
    workbook.save(temporary_path)
    temporary_path.replace(report_path)
    return report_path
